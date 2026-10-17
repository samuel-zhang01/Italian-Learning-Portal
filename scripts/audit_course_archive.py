from __future__ import annotations

import argparse
import json
import os
import re
import sys
import zipfile
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET

PORTAL = Path(__file__).resolve().parents[1]
MATERIALS = PORTAL / "materials"
OUT = PORTAL / "data" / "incoming"
SKIP_PARTS = {"__MACOSX", "__pycache__", "node_modules"}
TEXT_LIMIT = 120_000


def clean(text: str) -> str:
    text = text.replace("\x00", " ")
    return re.sub(r"\s+", " ", text).strip()


class TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def html_text(path: Path) -> str:
    parser = TextExtractor()
    parser.feed(path.read_text(errors="ignore"))
    return clean(" ".join(parser.parts))


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    return clean("\n".join((page.extract_text() or "") for page in PdfReader(path).pages))


def docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return clean("\n".join(parts))


def pptx_text(path: Path) -> str:
    texts: list[str] = []
    with zipfile.ZipFile(path) as zf:
        slides = sorted(n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
        for slide in slides:
            root = ET.fromstring(zf.read(slide))
            texts.extend(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
    return clean("\n".join(texts))


def xlsx_text(path: Path) -> str:
    texts: list[str] = []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name == "xl/sharedStrings.xml" or re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                root = ET.fromstring(zf.read(name))
                texts.extend(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
    return clean("\n".join(texts))


def eml_text(path: Path) -> tuple[str, dict[str, str]]:
    msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    bodies: list[str] = []
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_disposition() == "attachment":
                continue
            if part.get_content_type() == "text/plain":
                try:
                    bodies.append(part.get_content())
                except Exception:
                    pass
    elif msg.get_content_type() == "text/plain":
        bodies.append(msg.get_content())
    meta = {
        "subject": str(msg.get("subject", "")),
        "date": str(msg.get("date", "")),
        "from": str(msg.get("from", "")),
    }
    return clean("\n".join(bodies)), meta


def extract(path: Path) -> tuple[str, dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return html_text(path), {}
    if suffix == ".pdf":
        return pdf_text(path), {}
    if suffix == ".docx":
        return docx_text(path), {}
    if suffix == ".pptx":
        return pptx_text(path), {}
    if suffix == ".xlsx":
        return xlsx_text(path), {}
    if suffix == ".eml":
        return eml_text(path)
    if suffix in {".md", ".txt"}:
        return clean(path.read_text(errors="ignore")), {}
    return "", {}


def material_files(source: Path):
    # Prune before walking; do not follow links or index hidden credential files.
    for directory, folders, files in os.walk(source, followlinks=False):
        folders[:] = sorted(name for name in folders if not name.startswith(".")
                            and name not in SKIP_PARTS
                            and not (Path(directory) / name).is_symlink())
        for name in sorted(files):
            path = Path(directory) / name
            if not name.startswith(".") and not path.is_symlink() and path.is_file():
                yield path


def main() -> None:
    parser = argparse.ArgumentParser(description="Inventory local study materials without importing them into the app.")
    parser.add_argument("--source", type=Path, default=MATERIALS / "inbox",
                        help="Study folder under materials/ (default: materials/inbox)")
    args = parser.parse_args()
    source = args.source.resolve()
    if MATERIALS.is_symlink() or not source.is_relative_to(MATERIALS):
        parser.error("The source must be inside this project's materials/ folder.")
    if not source.is_dir():
        parser.error(f"Source folder does not exist: {source}")
    paths = sorted(material_files(source))
    if not paths:
        parser.error("No materials found. Add files to materials/inbox/ first; existing reports were left untouched.")

    inventory: list[dict[str, object]] = []
    corpus: list[dict[str, object]] = []
    for path in paths:
        rel = path.relative_to(source).as_posix()
        record: dict[str, object] = {
            "path": rel,
            "extension": path.suffix.lower(),
            "bytes": path.stat().st_size,
        }
        try:
            text, meta = extract(path)
            record.update(meta)
            record["characters_extracted"] = len(text)
            if text:
                corpus.append({"path": rel, "text": text[:TEXT_LIMIT], **meta})
        except Exception as exc:
            record["error"] = f"{type(exc).__name__}: {exc}"
        inventory.append(record)

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "archive-inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "archive-corpus.json").write_text(json.dumps(corpus, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = {
        "files": len(inventory),
        "extractable_files": len(corpus),
        "extensions": {},
        "errors": [x for x in inventory if "error" in x],
    }
    for item in inventory:
        ext = str(item["extension"] or "[none]")
        summary["extensions"][ext] = summary["extensions"].get(ext, 0) + 1
    (OUT / "archive-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    sys.exit(main())
